"""Filtrelenmiş kesinti listesini Excel / PDF / Word / CSV / JSON olarak üretir."""
import csv
import io
import json
import os
from datetime import datetime

from . import config
from .queries import filter_summary, fmt_dt, fmt_hours

TITLE = "MUSKİ Su Kesintileri Kaydı"

# (başlık, alan/fonksiyon, excel genişliği)
FULL_COLUMNS = [
    ("Kayıt No", "id", 9),
    ("Kayıt Türü", "kaynak_label", 12),
    ("MUSKİ Kesinti ID", "muski_id", 14),
    ("İş Emri No", "ays_is_emri_id", 12),
    ("İl", "il", 8),
    ("İlçe", "ilce", 14),
    ("İlçe ID", "ilce_id", 8),
    ("İlçe UAVT", "ilce_uavt", 9),
    ("Kesinti Tarihi", lambda o: fmt_dt(o["kesinti_tarihi"], True), 19),
    ("Tahmini Süre (saat)", "tahmini_sure_saat", 10),
    ("Tahmini Bitiş", lambda o: fmt_dt(o["tahmini_bitis"]), 17),
    ("Durum", "durum_label", 20),
    ("Sitedeki Durum", "site_durum_text", 14),
    ("İş Emri Durumu", "is_emri_durum_adi", 14),
    ("Kesinti Türü", "kesinti_turu", 11),
    ("Tür Etiketi", "tur_etiketi", 15),
    ("Etkilenen Mahalleler", "mahalleler_text", 50),
    ("Kesinti Yeri", "ariza_yeri", 28),
    ("Sebep", "sebep", 30),
    ("İlk Bildirilen Bitiş (SMS)", lambda o: fmt_dt(o["ilk_ilan_bitis"]), 17),
    ("Sahaya Gönderildi", "sahaya_gonderildi", 9),
    ("Boru Cinsi", "boru_cinsi", 10),
    ("İlk Görülme", lambda o: fmt_dt(o["first_seen_at"], True), 19),
    ("Son Görülme", lambda o: fmt_dt(o["last_seen_at"], True), 19),
    ("Son 'Devam' Görülme", lambda o: fmt_dt(o["last_active_seen_at"], True), 19),
    ("'Tamamlandı' Görülme", lambda o: fmt_dt(o["completed_seen_at"], True), 19),
    ("Siteden Kalkma", lambda o: fmt_dt(o["gone_at"], True), 19),
    ("Bitiş (en erken)", lambda o: fmt_dt(o["bitis_en_erken"]), 17),
    ("Bitiş (en geç)", lambda o: fmt_dt(o["bitis_en_gec"]), 17),
    ("Süre en az (saat)", "sure_en_az_saat", 10),
    ("Süre en fazla (saat)", "sure_en_fazla_saat", 10),
    ("Tahmin Aşıldı", lambda o: "Evet" if o["tahmin_asildi"] else "", 8),
    ("Not", "notes", 30),
]

SHORT_COLUMNS = [
    ("No", lambda o: str(o["id"]), 0.5),
    ("Tür", lambda o: {"web": "Web", "mesaj": "Mesaj"}.get(o["source"], o["source"]), 0.6),
    ("MUSKİ ID /\nİş Emri", lambda o: "\n".join(str(x) for x in (o["muski_id"], o["ays_is_emri_id"]) if x), 1.0),
    ("İlçe", "ilce", 1.1),
    ("Kesinti Tarihi", lambda o: fmt_dt(o["kesinti_tarihi"]), 1.2),
    ("Tahmini Süre /\nBitiş", lambda o: "\n".join(
        x for x in (fmt_hours(o["tahmini_sure_saat"]), fmt_dt(o["tahmini_bitis"])) if x), 1.1),
    ("Durum", "durum_label", 1.1),
    ("Kesinti Yeri", "ariza_yeri", 1.4),
    ("Etkilenen Mahalleler", lambda o: (o["mahalleler_text"] or "")
        + (f"\nSebep: {o['sebep']}" if o.get("sebep") else ""), 3.0),
    ("Bitiş (en geç)", lambda o: fmt_dt(o["bitis_en_gec"]), 1.2),
    ("Süre (en fazla)", lambda o: fmt_hours(o["sure_en_fazla_saat"]), 0.9),
]


def value(o, getter):
    v = getter(o) if callable(getter) else o.get(getter)
    return "" if v is None else v


def stamp():
    return datetime.now(config.TZ)


def filename(ext):
    return f"muski-kesintiler-{stamp():%Y%m%d-%H%M}.{ext}"


# --------------------------------------------------------------------------- Excel

def to_xlsx(rows, nbs, f) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Kesintiler"
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    ws.append([c[0] for c in FULL_COLUMNS])
    for o in rows:
        ws.append([value(o, c[1]) for c in FULL_COLUMNS])
    for i, c in enumerate(FULL_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = c[2]
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Mahalleler")
    ws2.append(["Kayıt No", "MUSKİ Kesinti ID", "İlçe", "Kesinti Tarihi", "Durum", "Mahalle", "Mahalle ID", "Mahalle UAVT"])
    for o in rows:
        for n in nbs.get(o["id"], []):
            ws2.append([o["id"], o["muski_id"], o["ilce"], fmt_dt(o["kesinti_tarihi"], True), o["durum_label"],
                        n["mahalle"], n["mahalle_id"], n["uavt_kodu"]])
    for i, w in enumerate([9, 14, 14, 19, 20, 24, 11, 12], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
        ws2.cell(row=1, column=i).font, ws2.cell(row=1, column=i).fill = head_font, head_fill
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    ws3 = wb.create_sheet("Bilgi")
    for row in (("Rapor", TITLE), ("Oluşturulma", f"{stamp():%d.%m.%Y %H:%M:%S}"),
                ("Filtre", filter_summary(f)), ("Kayıt sayısı", len(rows)),
                ("Kaynak", config.PAGE_URL), ("Kaynak (JSON)", config.API_URL),
                ("Not", "Bitiş zamanları, sitenin 15 dakikada bir taranmasıyla tespit edilir: "
                        "'en erken' = son 'devam ediyor' görülme, 'en geç' = ilk 'tamamlandı' görülme "
                        "ya da siteden kalkma zamanı.")):
        ws3.append(row)
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- PDF

def _font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    reg = next((p for p in config.FONT_PATHS if p and os.path.exists(p)), None)
    bold = next((p for p in config.FONT_BOLD_PATHS if p and os.path.exists(p)), None)
    if not reg:
        return "Helvetica", "Helvetica-Bold"
    if "TRFont" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("TRFont", reg))
        pdfmetrics.registerFont(TTFont("TRFont-Bold", bold or reg))
    return "TRFont", "TRFont-Bold"


def to_pdf(rows, f) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    font, bold = _font()
    cell = ParagraphStyle("c", fontName=font, fontSize=7.5, leading=9)
    head = ParagraphStyle("h", fontName=bold, fontSize=7.5, leading=9, textColor=colors.white)
    title = ParagraphStyle("t", fontName=bold, fontSize=14, leading=18)
    info = ParagraphStyle("i", fontName=font, fontSize=8.5, leading=11)

    def P(text, style=cell):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)

    data = [[P(c[0], head) for c in SHORT_COLUMNS]]
    for o in rows:
        data.append([P(value(o, c[1])) for c in SHORT_COLUMNS])
    total_w = landscape(A4)[0] - 2 * cm
    unit = total_w / sum(c[2] for c in SHORT_COLUMNS)
    table = LongTable(data, colWidths=[c[2] * unit for c in SHORT_COLUMNS], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#999999")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
    ]))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 7)
        canvas.drawString(cm, 0.6 * cm, f"{TITLE} — {stamp():%d.%m.%Y %H:%M} — Kaynak: {config.PAGE_URL}")
        canvas.drawRightString(landscape(A4)[0] - cm, 0.6 * cm, f"Sayfa {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=cm, rightMargin=cm,
                            topMargin=cm, bottomMargin=1.2 * cm, title=TITLE)
    story = [P(TITLE, title), Spacer(1, 4),
             P(f"Oluşturulma: {stamp():%d.%m.%Y %H:%M}   •   Kayıt sayısı: {len(rows)}", info),
             P(f"Filtre: {filter_summary(f)}", info), Spacer(1, 8), table]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# --------------------------------------------------------------------------- Word

def to_docx(rows, f) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.2))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9)

    doc.add_heading(TITLE, level=1)
    doc.add_paragraph(f"Oluşturulma: {stamp():%d.%m.%Y %H:%M}   •   Kayıt sayısı: {len(rows)}")
    doc.add_paragraph(f"Filtre: {filter_summary(f)}")
    doc.add_paragraph(f"Kaynak: {config.PAGE_URL}")

    table = doc.add_table(rows=1, cols=len(SHORT_COLUMNS))
    table.style = "Table Grid"
    total = sum(c[2] for c in SHORT_COLUMNS)
    widths = [Cm(27.3 * c[2] / total) for c in SHORT_COLUMNS]

    def set_cell(cell, text, bold=False, white=False):
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(text))
        run.font.size = Pt(8)
        run.bold = bold
        if white:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for i, c in enumerate(SHORT_COLUMNS):
        cell = table.rows[0].cells[i]
        set_cell(cell, c[0], bold=True, white=True)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "1F4E79")
        cell._tc.get_or_add_tcPr().append(shd)
    # başlık satırı her sayfada tekrar etsin
    trPr = table.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)

    for o in rows:
        cells = table.add_row().cells
        for i, c in enumerate(SHORT_COLUMNS):
            set_cell(cells[i], value(o, c[1]))
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- CSV / JSON

def to_csv(rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([c[0] for c in FULL_COLUMNS])
    for o in rows:
        w.writerow([value(o, c[1]) for c in FULL_COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # Excel Türkçe karakterleri doğru açsın


def to_json(rows, nbs, f) -> bytes:
    skip = {"raw_html_card"}
    out = []
    for o in rows:
        d = {k: v for k, v in o.items() if k not in skip}
        if d.get("raw_api_json"):
            d["raw_api_json"] = json.loads(d["raw_api_json"])
        d["mahalle_detay"] = nbs.get(o["id"], [])
        out.append(d)
    return json.dumps({"rapor": TITLE, "olusturulma": stamp().isoformat(timespec="seconds"),
                       "filtre": filter_summary(f), "kayit_sayisi": len(out), "kesintiler": out},
                      ensure_ascii=False, indent=2).encode("utf-8")
